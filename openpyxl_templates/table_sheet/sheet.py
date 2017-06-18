from collections import Counter
from itertools import chain

from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from openpyxl_templates.exceptions import OpenpyxlTemplateException
from openpyxl_templates.table_sheet.columns import TableColumn
from openpyxl_templates.templated_sheet import TemplatedSheet
from openpyxl_templates.utils import Typed, MAX_COLUMN_INDEX


class ColumnHeadersNotUnique(OpenpyxlTemplateException):
    def __init__(self, columns):
        counter = Counter(column.header for column in columns)
        super().__init__("headers '%s' has been declared more then once in the same TableSheet" % tuple(
            header
            for (header, count)
            in counter.items()
            if count > 1
        ))


class TempleteStyleNotFound(OpenpyxlTemplateException):
    def __init__(self, missing_style_name, style_set):
        super().__init__(
            "The style '%s' has not been declared. Avaliable styles are: %s)"
            % (missing_style_name, style_set.names)
        )


class NoTableColumns(OpenpyxlTemplateException):
    def __init__(self, table_sheet):
        super().__init__(
            "The TableSheet '%s' has no columns. Declare atleast one."
            % table_sheet.sheetname
        )


class TableSheet(TemplatedSheet):
    item_class = TableColumn

    title_style = Typed("title_style", expected_type=str, value="Title")
    description_style = Typed("description_style", expected_type=str, value="Description")

    format_as_table = Typed("format_as_header", expected_type=bool, value=True)
    freeze_header = Typed("freeze_header", expected_type=bool, value=True)
    hide_excess_columns = Typed("hide_excess_columns", expected_type=bool, value=True)

    _first_data_cell = None
    _last_data_cell = None
    _first_header_cell = None
    _last_header_cell = None

    def __init__(self, sheetname=None, active=None):
        super().__init__(sheetname=sheetname, active=active)

        self.columns = list(self._items.values())

        for index, column in enumerate(self.columns):
            column.column_index = index + 1  # Start as 1

        self._validate()

    def _validate(self):
        self._check_atleast_one_column()
        self._check_unique_column_headers()

    def _check_atleast_one_column(self):
        if not self.columns:
            raise NoTableColumns(self)

    def _check_unique_column_headers(self):
        if len(set(column.header for column in self.columns)) < len(self.columns):
            raise ColumnHeadersNotUnique(self.columns)

    def write(self, title=None, description=None, objects=None):
        worksheet = self.worksheet

        self.prepare_worksheet(worksheet)
        self.write_title(worksheet, title)
        self.write_description(worksheet, description)
        self.write_headers(worksheet)
        self.write_rows(worksheet, objects)
        self.post_process_worksheet(worksheet)

    def prepare_worksheet(self, worksheet):
        for column in self.columns:
            column.prepare_worksheet(worksheet)

        # Register styles
        style_names = set(chain(
            (self.title_style, self.description_style),
            *((column.row_style, column.header_style) for column in self.columns)
        ))

        existing_names = set(self.workbook.named_styles)

        for name in style_names:
            if name in existing_names:
                continue

            if name not in self.workbook.template_styles:
                raise TempleteStyleNotFound(name, self.workbook.template_styles)

            self.workbook.add_named_style(self.workbook.template_styles[name])

    def write_title(self, worksheet, title=None):
        if not title:
            return

        title = WriteOnlyCell(ws=worksheet, value=title)
        title.style = self.title_style

        worksheet.append((title,))

    def write_description(self, worksheet, description=None):
        if not description:
            return

        description = WriteOnlyCell(ws=worksheet, value=description)
        description.style = self.description_style

        worksheet.append((description,))

    def write_headers(self, worksheet):
        headers = tuple(
            column.create_header(worksheet)
            for column in self.columns
        )

        self.worksheet.append(headers)

        self._first_header_cell = headers[0]
        self._last_header_cell = headers[-1]

    def write_rows(self, worksheet, objects=None):
        self._first_data_cell = None
        cells = None
        for obj in objects:
            cells = tuple(column.create_cell(worksheet, obj) for column in self.columns)
            worksheet.append(cells)

            if not self._first_data_cell:
                self._first_data_cell = cells[0]

            for cell, column in zip(cells, self.columns):
                column.post_process_cell(worksheet, cell)

        if cells:
            self._last_data_cell = cells[-1]

    def post_process_worksheet(self, worksheet):
        for column in self.columns:
            column.post_process_worksheet(worksheet)

        if self.active:
            self.activate()

        if self.format_as_table:
            worksheet.add_table(
                Table(
                    ref="%s:%s" % (
                        self._first_header_cell.coordinate,
                        self._last_data_cell.coordinate if self._last_data_cell else self._last_header_cell.coordinate
                    ),
                    displayName=self.sheetname,
                )
            )

        if self.freeze_header:
            worksheet.freeze_panes = self._first_data_cell or self._first_header_cell

        if self.hide_excess_columns:
            for i in range(len(self.columns) + 1, MAX_COLUMN_INDEX + 1):
                worksheet.column_dimensions[get_column_letter(i)].hidden = True

    def read(self, exception_policy):
        pass
